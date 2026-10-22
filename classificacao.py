"""
Preenche Débito, Crédito e Histórico da planilha base.

- Saída do banco (despesa): débito na conta do fornecedor, crédito na conta do banco.
- Entrada no banco (receita): débito na conta do banco, crédito na conta do cliente.

A conta do fornecedor ou cliente vem da planilha matriz (descrição -> conta contábil),
procurando o nome que aparece no extrato. A conta de cada banco vem da tabela de contas
dos bancos (Banco/Conta da planilha base -> código). O que não dá para decidir fica em
branco, com o motivo na coluna "Conferir".
"""

import io
import re
import unicodedata

import pandas as pd

HIST_SAIDA = "VR RF A SAIDA DO BANCO - "
HIST_ENTRADA = "VR RF A ENTRADA NO BANCO - "
NAO_IMPORTAR = "não importar"

# Palavras que não ajudam a reconhecer o fornecedor
_VAZIAS = {"ltda", "me", "mei", "epp", "eireli", "sa", "s", "a", "de", "da", "do", "das", "dos",
           "e", "em", "cia", "comercio", "servicos", "filial", "nf", "pgto", "pg", "ref"}

# O que vem antes do nome do fornecedor/cliente na descrição da matriz
_PREFIXOS_MATRIZ = re.compile(
    r"^(?:dep[oó]sito \w+ \([^)]*\)? ?-|adiantamento reserva -|pix recep[cç][aã]o -|"
    r"cch -|bnb - nf -|bnb -|nf -|transf entre contas -> sunbay \|)\s*", re.IGNORECASE)

# O que vem antes do nome da contraparte na descrição do extrato
_PREFIXOS_EXTRATO = re.compile(
    r"^(?:.*? - )?(?:rem:|remet\.|dest\.|fav:?)\s*|"
    r"^(?:pix enviado|pix recebido|boleto pago|pagto eletron cobranca|sispag|"
    r"transf cc para cc pj -|ted-credito em conta - \d+ \d+ \d+|ted [\d.]+)\s*", re.IGNORECASE)

# Lançamentos de banco com nome genérico: o texto do extrato e o nome procurado na matriz
_APELIDOS = [
    (r"^pix qrs\b", "pix recepcao"),
    (r"^rede\s+(visa|mast|elo|amex|hiper)", "rede"),
    (r"^(cielo|getnet|stone)\b", r"\1"),
    (r"^(tarifa|tar)\b", "tarifa bancaria"),
    (r"^iof\b", "iof"),
]

# Aplicação e resgate na conta corrente: a outra conta é a aplicação do mesmo banco
_APLICACAO = re.compile(r"^(apl|aplic|res aplic|resgate|rend pago|rentab|resg)\b", re.IGNORECASE)
_APLICACAO_ESPELHO = {"Aplicação", "Resgate"}   # linhas da conta de aplicação


def _sem_acento(texto):
    t = unicodedata.normalize("NFKD", str(texto))
    return "".join(c for c in t if not unicodedata.combining(c)).lower()


def _tokens(texto):
    palavras = re.findall(r"[a-z0-9]+", _sem_acento(texto))
    return [p for p in palavras if p not in _VAZIAS and not p.isdigit()]


def _conserta(texto):
    """Desfaz acento quebrado de planilha que passou por outra codificação
    ("SERVIÃ‡O" -> "SERVIÇO")."""
    if "Ã" not in texto and "Â" not in texto:
        return texto

    def byte(c):   # cp1252 cobre "“" (0x93); os bytes que ele não tem (0x81, 0x8d) vêm crus
        try:
            return c.encode("cp1252")
        except UnicodeEncodeError:
            return c.encode("latin-1") if ord(c) < 256 else c.encode("utf-8")
    try:
        return b"".join(byte(c) for c in texto).decode("utf-8")
    except UnicodeDecodeError:
        return texto


def _tabelas_da_matriz(arquivo):
    """Pares (descrição, conta) da primeira aba. A matriz pode ter uma tabela só (descrição
    e conta nas colunas A e B) ou várias lado a lado, cada uma com cabeçalho próprio:
    PARTICIPANTE | CONTA, BANCO | CONTA, DESCRIÇÃO DA RECEITA | CONTA, DESCRIÇÃO DA
    DESPESA | CONTA. Devolve [(tipo, DataFrame com Descrição e Conta)], com tipo
    "banco", "receita", "despesa" ou "" (vale para os dois sentidos)."""
    bruto = pd.read_excel(arquivo, header=None, dtype=str).fillna("")
    bruto = bruto.map(lambda v: _conserta(str(v).strip()))
    linha, pares = _cabecalho_da_matriz(bruto)
    if pares:
        tabelas = []
        for tipo, j in pares:
            t = bruto.iloc[linha + 1:, [j - 1, j]]
            t.columns = ["Descrição", "Conta"]
            tabelas.append((tipo, t))
        return tabelas
    t = bruto.iloc[:, :2].copy()
    t.columns = ["Descrição", "Conta"]
    return [(None, t)]


def _cabecalho_da_matriz(bruto):
    """(linha do cabeçalho, [(tipo, coluna da CONTA)]) da matriz com várias tabelas, ou
    (None, []) quando é uma tabela só."""
    for linha in range(min(len(bruto), 10)):
        cab = [_sem_acento(v) for v in bruto.iloc[linha]]
        pares = [j for j in range(1, len(cab)) if cab[j].startswith("conta") and cab[j - 1]]
        if len(pares) > 1:
            return linha, [("banco" if cab[j - 1].startswith("banco") else
                            "receita" if "receita" in cab[j - 1] else
                            "despesa" if "despesa" in cab[j - 1] else "", j) for j in pares]
    return None, []


def ler_matriz(arquivo):
    """Planilha matriz -> DataFrame com Descrição, Conta, Nome (tokens do nome) e
    Recebimento (True: só entradas; False: só saídas; None: serve para as duas)."""
    return matriz_da_tabela(tabela_da_matriz(arquivo))[0]


def bancos_na_matriz(arquivo):
    """{banco: conta} da tabela BANCO | CONTA da matriz, quando ela existe."""
    for tipo, t in _tabelas_da_matriz(arquivo):
        if tipo == "banco":
            t = t[(t["Descrição"] != "") & (t["Conta"] != "")]
            return dict(zip(t["Descrição"], t["Conta"].str.replace(r"\.0$", "", regex=True)))
    return {}


def parece_matriz(nome, conteudo):
    """O arquivo é a planilha matriz? Pelo nome ou pelo cabeçalho "Conta Contábil"."""
    nome = _sem_acento(nome)
    if not nome.endswith((".xlsx", ".xls")):
        return False
    if "matriz" in nome:
        return True
    try:
        topo = pd.read_excel(io.BytesIO(conteudo), header=None, dtype=str, nrows=10)
    except Exception:
        return False
    return topo.shape[1] >= 2 and topo.iloc[:, 1].fillna("").map(_sem_acento) \
        .str.startswith("conta cont").any()


def parece_planilha_base(nome, conteudo):
    """O arquivo é uma planilha base baixada deste app (não é extrato nem matriz)?"""
    if not _sem_acento(nome).endswith((".xlsx", ".xls")):
        return False
    try:
        topo = pd.read_excel(io.BytesIO(conteudo), header=None, dtype=str, nrows=1)
    except Exception:
        return False
    return {"Banco/Conta", "Débito", "Crédito", "Histórico"} <= set(topo.iloc[0].fillna(""))


# Linhas da matriz que são de dinheiro entrando (cliente pagando, depósito, cartão)
_RECEBIMENTO = re.compile(r"(?i)^(dep[oó]sito|pix recep|credito\b|adiantamento reserva|"
                          r"recupera)")

# Primeira parte que não é nome de ninguém ("CREDITO - REDE": o nome é REDE)
_GENERICOS = {"credito", "debito", "pagamento", "deposito", "recebimento", "banco", "bradesco",
              "itau", "bb", "bnb", "santander", "caixa", "sicredi"}


def _nome_na_matriz(descricao):
    d = _PREFIXOS_MATRIZ.sub("", descricao.strip())
    # "FORNECEDOR - CATEGORIA": o nome é a primeira parte que não é genérica
    partes = [x for x in d.split(" - ") if _tokens(x)]
    for parte in partes:
        nome = tuple(t for t in _tokens(parte) if not t.isdigit())
        if nome and not set(nome) <= _GENERICOS:
            return nome
    return tuple(_tokens(d))


def _nomes_no_extrato(descricao):
    """Nomes a procurar na matriz, do mais provável ao menos. O nome costuma ser a última
    parte ("TRANSFERENCIA PIX - REM: FULANO"); nas planilhas do cliente o histórico junta
    várias colunas ("CATEGORIA - FORNECEDOR - CNPJ"), e aí vale tentar cada parte."""
    d = re.sub(r"\s*\d{2}/\d{2}$", "", descricao.strip())     # data do Pix no fim
    d = re.sub(r"\s+\d{2,3}/\d{2,3}$", "", d)                   # parcela do cartão
    for padrao, apelido in _APELIDOS:
        m = re.search(padrao, _sem_acento(d))
        if m:
            return [tuple(_tokens(m.expand(apelido)))]
    partes = d.split(" - ")
    nomes = [tuple(_tokens(_PREFIXOS_EXTRATO.sub("", p))) for p in reversed(partes)]
    if len(partes) > 1:
        nomes.append(tuple(_tokens(d)))
    return [n for n in dict.fromkeys(nomes) if n]


def _nome_no_extrato(descricao):
    nomes = _nomes_no_extrato(descricao)
    return nomes[0] if nomes else ()


def _casa(a, b, cortado):
    """Palavra do extrato casa com a da matriz. Só a última palavra de um nome pode estar
    cortada ("ORINTER VIAGENS E TUR"), e aí basta ser o começo da outra."""
    if a == b:
        return True
    return cortado and min(len(a), len(b)) >= 3 and (a.startswith(b) or b.startswith(a))


def _pontua(extrato, matriz):
    if not extrato or not matriz:
        return 0.0
    junto_e, junto_m = "".join(extrato), "".join(matriz)
    if min(len(junto_e), len(junto_m)) >= 6 and (junto_e.startswith(junto_m)
                                                  or junto_m.startswith(junto_e)):
        return 1.0   # "DIGITALOCEAN" x "DIGITAL OCEAN"
    if not _casa(extrato[0], matriz[0], len(extrato) == 1 or len(matriz) == 1):
        return 0.0
    curto, longo = (extrato, matriz) if len(extrato) <= len(matriz) else (matriz, extrato)
    iguais = sum(any(_casa(p, q, i == len(curto) - 1 or j == len(longo) - 1)
                     for j, q in enumerate(longo)) for i, p in enumerate(curto))
    # Nome cortado no extrato ("ORINTER VIAGENS E TUR" x "ORINTER TOUR E TRAVEL"): a
    # primeira palavra, se for longa e igual, já identifica
    primeira = 0.75 if extrato[0] == matriz[0] and len(extrato[0]) >= 5 else 0
    return max(iguais / len(curto), primeira)


_FINANCEIRA = re.compile(
    r"\b(banco|bancoob|sicoob|sicredi|unicred|cresol|cooperativ[ao] de credito|caixa economica|"
    r"financeira|financiamento|consorcio|asaas|pagseguro|mercado pago|stone|cielo|getnet|nubank|"
    r"itau|bradesco|santander|btg|safra|banrisul)\b")
_EMPRESA = re.compile(r"\d{11,14}|\bs\s*[./]\s*a\b|\bltda\b")


def eh_instituicao_financeira(nome):
    """O nome (da matriz) é de um banco, cooperativa de crédito, financeira ou similar? Só
    nomes de empresa (com CNPJ, S.A. ou LTDA): "TAXAS BANCÁRIAS" e "CARTÃO SICOOB" são
    categorias, não instituições."""
    n = _sem_acento(str(nome))
    return bool(_FINANCEIRA.search(n) and _EMPRESA.search(n))


def procurar_na_matriz(descricao, matriz, minimo=0.75, entrada=None):
    """(conta, descrição da matriz, aviso) para a descrição do extrato.

    entrada: True/False para preferir as linhas da matriz do mesmo sentido (o mesmo cliente
    aparece como "Depósito ... - EUROPLUS" e o fornecedor como "EUROPLUS - ...")."""
    nomes = _nomes_no_extrato(descricao)
    if not nomes:
        return None, "", "sem nome do fornecedor no extrato"
    # Receita só vale para entradas e Despesa só para saídas; Participante, para as duas
    if entrada is not None and "Recebimento" in matriz:
        validas = [i for i, r in enumerate(matriz["Recebimento"]) if r is None or bool(r) == entrada]
    else:
        validas = range(len(matriz))
    for k, nome in enumerate(nomes):
        notas = [(_pontua(nome, matriz.at[i, "Nome"]), i) for i in validas]
        maior = max((s for s, _ in notas), default=0)
        # As outras partes só valem se casarem por inteiro ("TRANSF CC PARA CC PJ" não é
        # "TRANSF ENTRE CONTAS")
        if maior >= (minimo if k == 0 else 1.0):
            break
    else:
        return None, "", "não achei na matriz"
    melhores = [i for s, i in notas if s == maior]   # só os que casaram melhor votam
    iguais = [i for i in melhores if matriz.at[i, "Nome"] == nome]
    melhores = iguais or melhores
    if entrada is not None and "Recebimento" in matriz:
        # A linha do mesmo sentido ganha da que vale para os dois ("Depósito ... - QUANTA")
        proprias = [i for i in melhores if matriz.at[i, "Recebimento"] is not None]
        melhores = proprias or melhores
    contas = matriz.loc[melhores, "Conta"]
    votos = contas.value_counts()
    conta = votos.index[0]
    exemplo = matriz.loc[contas[contas == conta].index[0], "Descrição"]
    aviso = ""
    if len(votos) > 1:
        aviso = "matriz com contas diferentes: " + ", ".join(f"{c} ({n}x)" for c, n in
                                                              votos.items())
    return conta, exemplo, aviso


def ler_contas_bancos(arquivo):
    """Planilha (xlsx ou csv) com Banco/Conta e o código contábil -> {Banco/Conta: código}."""
    nome = str(getattr(arquivo, "name", arquivo)).lower()
    if nome.endswith(".csv"):
        t = pd.read_csv(arquivo, sep=None, engine="python", dtype=str)
    else:
        t = pd.read_excel(arquivo, dtype=str)
    t = t.iloc[:, :2].dropna()
    return {str(c).strip(): str(v).strip().removesuffix(".0") for c, v in t.itertuples(index=False)}


def codigos_na_matriz(arquivo):
    """Códigos das contas guardados na matriz: uma aba com as colunas Banco/Conta e Código
    (Banco/Conta escrita como aparece na planilha base). {} se a matriz não tiver."""
    for aba in pd.read_excel(arquivo, sheet_name=None, header=None, dtype=str).values():
        aba = aba.fillna("").map(lambda v: str(v).strip())
        for i, linha in aba.iterrows():
            celulas = list(linha)
            if "Banco/Conta" in celulas and "Código" in celulas:
                c_conta, c_codigo = celulas.index("Banco/Conta"), celulas.index("Código")
                abaixo = aba.iloc[i + 1:]
                return {conta: codigo.removesuffix(".0") for conta, codigo in
                        zip(abaixo[c_conta], abaixo[c_codigo]) if conta and codigo}
    return {}


def _codigo_do_banco(conta, contas_bancos):
    return str(contas_bancos.get(conta, "") or "").strip()


def _aplicacao_do_banco(conta, contas_bancos):
    """Código da aplicação ligada à conta corrente (mesmo banco, mesmo número de conta)."""
    def numero(nome):
        m = re.search(r"Cc (\S+)$", nome)
        return re.sub(r"\D", "", m.group(1)).lstrip("0") if m else None

    proprio = numero(conta)
    if not proprio:
        return ""
    for outra, codigo in contas_bancos.items():
        if outra != conta and numero(outra) == proprio and codigo and "Cartão" not in outra:
            return str(codigo)
    return ""


def classificar(base, matriz, contas_bancos):
    """Base com Débito, Crédito e Histórico preenchidos, mais "Conta na matriz" e "Conferir".

    contas_bancos: {Banco/Conta como está na base: código contábil}.
    """
    b = base.copy()
    debitos, creditos, historicos, origens, conferir = [], [], [], [], []
    for conta_banco, descricao, valor in zip(b["Banco/Conta"], b["Descrição"], b["Valor"]):
        banco = _codigo_do_banco(conta_banco, contas_bancos)
        avisos = []
        outra, origem = None, ""
        if valor == 0:
            debitos.append(""), creditos.append(""), historicos.append("")
            origens.append(""), conferir.append("sem movimento no período")
            continue
        if descricao in _APLICACAO_ESPELHO and "Aplic" in conta_banco or \
                descricao in _APLICACAO_ESPELHO and any(k in conta_banco for k in
                                                         ("Invest", "CDB", "Aplic")):
            debitos.append(""), creditos.append(""), historicos.append("")
            origens.append(NAO_IMPORTAR)
            conferir.append("já lançado pela conta corrente (aplicação/resgate)")
            continue
        if _APLICACAO.match(descricao) and "Cartão" not in conta_banco:
            outra = _aplicacao_do_banco(conta_banco, contas_bancos)
            origem = "aplicação do mesmo banco"
            if not outra:
                avisos.append("falta o código da aplicação deste banco")
        if not outra and not (origem and avisos):
            outra, origem, aviso = procurar_na_matriz(descricao, matriz, entrada=valor > 0)
            if aviso:
                avisos.append(aviso)
            if outra and eh_instituicao_financeira(origem):
                # Banco ou financeira como fornecedor/cliente: a conta depende do que é
                avisos.append("instituição financeira: confira se é " + (
                    "rendimento, resgate ou transferência" if valor > 0
                    else "despesa bancária, empréstimo ou consórcio"))
                outra, origem = None, f"(não usada) {origem}"
        if outra and str(outra).lower() == NAO_IMPORTAR:
            debitos.append(""), creditos.append(""), historicos.append("")
            origens.append(origem), conferir.append("matriz diz para não importar")
            continue
        if not banco:
            avisos.append("falta o código desta conta (banco ou cartão)")
        outra = outra or ""
        if valor < 0:
            debitos.append(outra), creditos.append(banco)
            historicos.append(HIST_SAIDA + descricao)
        else:
            debitos.append(banco), creditos.append(outra)
            historicos.append(HIST_ENTRADA + descricao)
        origens.append(origem)
        conferir.append("; ".join(avisos))
    b["Débito"], b["Crédito"], b["Histórico"] = debitos, creditos, historicos
    b["Conta na matriz"], b["Conferir"] = origens, conferir
    return b


def nome_para_matriz(descricao):
    """Texto a gravar numa linha nova da matriz: a parte do histórico com o nome (a última,
    sem prefixos como "PIX ENVIADO" e sem CNPJ ou data)."""
    d = re.sub(r"\s*\d{2}/\d{2}$", "", str(descricao).strip())
    d = re.sub(r"\s+\d{2,3}/\d{2,3}$", "", d)
    for parte in reversed(d.split(" - ")):
        parte = _PREFIXOS_EXTRATO.sub("", parte).strip(" -:")
        if parte and not re.fullmatch(r"[\d./-]+", parte):
            return parte
    return d


def linhas_novas_da_matriz(classificada, corrigida, matriz):
    """Lançamentos em que a conta do fornecedor/cliente foi corrigida ou preenchida à mão e
    que a matriz ainda não conhece: DataFrame com Descrição, Conta e Entrada (bool)."""
    linhas = []
    for i in corrigida.index:
        entrada = corrigida.at[i, "Valor"] > 0
        lado = "Crédito" if entrada else "Débito"
        conta = str(corrigida.at[i, lado]).strip()
        if not conta or conta == str(classificada.at[i, lado]).strip():
            continue
        linhas.append({"Descrição": nome_para_matriz(corrigida.at[i, "Descrição"]),
                       "Conta": conta, "Entrada": bool(entrada)})
    novas = pd.DataFrame(linhas, columns=["Descrição", "Conta", "Entrada"])
    novas = novas.drop_duplicates(["Descrição", "Entrada"], keep="last")
    if matriz is not None and len(novas):
        ja = {(_sem_acento(d), r) for d, r in zip(matriz["Descrição"], matriz["Recebimento"])}
        novas = novas[[(_sem_acento(d), e) not in ja and (_sem_acento(d), None) not in ja
                       for d, e in zip(novas["Descrição"], novas["Entrada"])]]
    return novas.reset_index(drop=True)


def matriz_mais(matriz, novas):
    """A matriz lida mais as linhas novas (ainda não gravadas no arquivo)."""
    extra = pd.DataFrame({"Descrição": novas["Descrição"], "Conta": novas["Conta"].astype(str),
                          "Recebimento": novas["Entrada"].astype(object)})
    extra["Nome"] = extra["Descrição"].map(_nome_na_matriz)
    return pd.concat([matriz, extra], ignore_index=True)


# Matriz como uma tabela só, para consultar e editar no app
TABELAS = ["Participante", "Receita", "Despesa", "Banco"]
_NOME_TABELA = {"": "Participante", "receita": "Receita", "despesa": "Despesa", "banco": "Banco"}


def tabela_da_matriz(arquivo):
    """Todas as linhas da matriz: DataFrame com Tabela, Descrição e Conta."""
    partes = []
    for tipo, t in _tabelas_da_matriz(arquivo):
        t = t.copy()
        t["Conta"] = t["Conta"].str.replace(r"\.0$", "", regex=True)
        cabecalho = t["Conta"].str.lower().str.startswith("conta")
        t = t[~cabecalho & (t["Descrição"] != "") & (t["Conta"] != "")]
        if tipo is None:   # tabela única: só "Depósito ...", "PIX RECEPÇÃO" etc. têm sentido
            t["Tabela"] = t["Descrição"].str.match(_RECEBIMENTO).map({True: "Receita",
                                                                       False: "Participante"})
        else:
            t["Tabela"] = _NOME_TABELA[tipo]
        partes.append(t[["Tabela", "Descrição", "Conta"]])
    return pd.concat(partes, ignore_index=True)


def matriz_da_tabela(tabela):
    """(matriz para classificar, {banco: conta}) a partir da tabela editada no app."""
    t = tabela.fillna("").astype(str).apply(lambda c: c.str.strip())
    t = t[(t["Descrição"] != "") & (t["Conta"] != "")]
    bancos = t[t["Tabela"] == "Banco"]
    m = t[t["Tabela"] != "Banco"].copy()
    m["Recebimento"] = m["Tabela"].map({"Receita": True, "Despesa": False}).astype(object)
    m["Recebimento"] = m["Recebimento"].where(m["Tabela"].isin(["Receita", "Despesa"]), None)
    m["Nome"] = m["Descrição"].map(_nome_na_matriz)
    return (m[["Descrição", "Conta", "Recebimento", "Nome"]].reset_index(drop=True),
            dict(zip(bancos["Descrição"], bancos["Conta"])))


def tabela_mais(tabela, novas):
    """A tabela da matriz mais as linhas novas que vieram das correções."""
    extra = pd.DataFrame({"Tabela": novas["Entrada"].map({True: "Receita", False: "Despesa"}),
                          "Descrição": novas["Descrição"], "Conta": novas["Conta"].astype(str)})
    return pd.concat([tabela, extra], ignore_index=True)


def matriz_para_excel(tabela, original=None):
    """A matriz em .xlsx, com as tabelas lado a lado (PARTICIPANTE | CONTA, BANCO | CONTA,
    DESCRIÇÃO DA RECEITA | CONTA, DESCRIÇÃO DA DESPESA | CONTA) na primeira aba. As outras
    abas do arquivo original (ex.: códigos dos bancos) continuam como estavam."""
    import openpyxl
    from openpyxl.styles import Font
    eh_xlsx = original is not None and original[:4] == b"PK\x03\x04"
    livro = openpyxl.load_workbook(io.BytesIO(original)) if eh_xlsx else openpyxl.Workbook()
    antiga = livro.worksheets[0]
    aba = livro.create_sheet(antiga.title or "Matriz", 0)
    livro.remove(antiga)
    t = tabela.fillna("").astype(str)
    t = t[(t["Descrição"].str.strip() != "") & (t["Conta"].str.strip() != "")]
    cabecalhos = {"Participante": "DESCRIÇÃO/PARTICIPANTE", "Banco": "BANCO",
                  "Receita": "DESCRIÇÃO DA RECEITA", "Despesa": "DESCRIÇÃO DA DESPESA"}
    for n, (nome, cabecalho) in enumerate(cabecalhos.items()):
        col = 1 + 3 * n
        aba.cell(1, col, cabecalho).font = Font(bold=True)
        aba.cell(1, col + 1, "CONTA CONTÁBIL").font = Font(bold=True)
        aba.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 45
        aba.column_dimensions[openpyxl.utils.get_column_letter(col + 1)].width = 16
        for linha, (_, r) in enumerate(t[t["Tabela"] == nome].iterrows(), start=2):
            conta = r["Conta"].strip()
            aba.cell(linha, col, r["Descrição"].strip())
            aba.cell(linha, col + 1, int(conta) if conta.isdigit() else conta)
    aba.freeze_panes = "A2"
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()
